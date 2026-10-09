"""Fluxo da ferramenta (o mesmo que a interface web executa), exportações e lote."""

import csv
import io
import json
from datetime import datetime, timedelta

import numpy as np
import openpyxl
import pytest

from mnp_nucleo import __version__
from mnp_nucleo.fantomas import gerar_dicom
from mnp_nucleo.gastrico import lote
from mnp_nucleo.gastrico.relatorio import AVISO
from mnp_nucleo.gastrico.sessao import Sessao

from .conftest import comando

INICIO = datetime(2026, 9, 21, 23, 30)


def dcm(minutos, vista, valor=100, tamanho=16, hora=True, **kw):
    img = np.full((tamanho, tamanho), valor)
    return gerar_dicom(img, INICIO + timedelta(minutes=minutos) if hora else None, vista, **kw)


def carregar(s, arquivos):
    for nome, dados in arquivos:
        s.adicionar_arquivo(nome, dados)
    return comando(s, acao="estado")["estado"]


def test_pareamento_cronologico_e_virada_da_meia_noite():
    s = Sessao()
    # fora de ordem e atravessando a meia-noite (23:30 → 03:30)
    e = carregar(
        s,
        [("c", dcm(240, "POST")), ("a", dcm(0, "ANT")), ("d", dcm(60, "ANT")), ("b", dcm(0, "POST")),
         ("e", dcm(60, "POST")), ("f", dcm(240, "ANT"))],
    )
    assert [t["minutos"] for t in e["tempos"]] == [0, 60, 240]
    nomes = {a["id"]: a["nome"] for a in e["arquivos"]}
    assert [(nomes[t["ant"]], nomes[t["post"]]) for t in e["tempos"]] == [("a", "b"), ("d", "e"), ("f", "c")]
    assert e["validacao"]["arquivos"] is None


def test_validacao_minimo_de_tempos():
    s = Sessao()
    e = carregar(s, [("a", dcm(0, "ANT")), ("b", dcm(0, "POST")), ("c", dcm(60, "ANT")), ("d", dcm(60, "POST"))])
    assert "pelo menos 3 tempos" in e["validacao"]["arquivos"]


def test_horario_ausente_e_edicao_manual_preservada():
    s = Sessao()
    e = carregar(
        s,
        [("a", dcm(0, "ANT")), ("b", dcm(0, "POST")), ("c", dcm(60, "ANT")), ("d", dcm(60, "POST")),
         ("x", dcm(0, "ANT", hora=False)), ("y", dcm(0, "POST", hora=False))],
    )
    t = e["tempos"][2]
    assert t["horario_ausente"] and t["minutos"] is None
    assert e["validacao"]["arquivos"] is not None
    assert comando(s, acao="definir_minutos", tempo=2, valor="180,5")["ok"]
    assert s.tempos[2].minutos == 180.5
    # adicionar outro arquivo não apaga a edição manual
    s.adicionar_arquivo("z", b"lixo" * 100)
    assert s.tempos[2].minutos == 180.5
    assert comando(s, acao="definir_minutos", tempo=2, valor="-3")["ok"] is False


def test_erro_de_arquivo_vira_mensagem():
    s = Sessao()
    e = json.loads(s.adicionar_arquivo("foto.jpg", b"\xff\xd8" * 200))["estado"]
    assert e["erros"] and e["erros"][0].startswith("foto.jpg:")
    assert comando(s, acao="limpar_erros")["estado"]["erros"] == []


def test_troca_de_vista_e_dimensoes_diferentes():
    s = Sessao()
    e = carregar(s, [("a", dcm(0, "ANT")), ("b", dcm(0, "ANT", tamanho=8))])
    assert len(e["tempos"]) == 2
    comando(s, acao="definir_vista", arquivo=e["arquivos"][1]["id"], vista="posterior")
    t = comando(s, acao="estado")["estado"]["tempos"][0]
    assert t["dimensoes_diferentes"] and "Dimensões" in t["bloqueio"]


def test_primeira_roi_copiada_e_ajuste_individual(sessao_pronta, fantoma):
    rois = sessao_pronta.rois
    assert all(rois[i] == rois[0] for i in range(4))
    # a segunda ROI desenhada não é copiada
    comando(sessao_pronta, acao="desenhar_roi", tempo=2, pontos=[[1, 1], [10, 1], [10, 10], [1, 10]])
    assert rois[2] != rois[0] and rois[3] == rois[0]
    comando(sessao_pronta, acao="copiar_roi", tempo=2)
    assert all(rois[i] == rois[2] for i in range(4))


def test_arraste_limitado_e_so_no_tempo(sessao_pronta):
    antes = list(sessao_pronta.rois[1])
    previa = comando(sessao_pronta, acao="arrastar_roi", tempo=1, dx=-500, dy=0, confirmar=False)
    assert min(p[0] for p in previa["roi"]) == 0
    assert sessao_pronta.rois[1] == antes  # prévia não grava
    comando(sessao_pronta, acao="arrastar_roi", tempo=1, dx=-500, dy=0, confirmar=True)
    assert min(p[0] for p in sessao_pronta.rois[1]) == 0
    assert sessao_pronta.rois[0] != sessao_pronta.rois[1]


def test_toca_roi(sessao_pronta, fantoma):
    x0, y0 = fantoma.roi[0]
    assert comando(sessao_pronta, acao="toca_roi", tempo=0, x=x0 + 3, y=y0 + 3)["toca"]
    assert not comando(sessao_pronta, acao="toca_roi", tempo=0, x=1, y=1)["toca"]


def test_troca_de_quadro_descarta_roi(fantoma):
    s = Sessao()
    for n, d in fantoma.multiframe.items():
        s.adicionar_arquivo(n, d)
    sug = comando(s, acao="estado")["estado"]["sugestoes"]
    for x in sug:
        comando(s, acao="aceitar_multiframe", arquivo=x["arquivo"])
    s.desenhar_roi(0, fantoma.roi)
    s.desenhar_eixo(0, fantoma.eixo)
    assert 1 in s.rois
    comando(s, acao="definir_quadro", arquivo=s.tempos[1].post, quadro=0)
    assert 1 not in s.rois and 0 in s.rois
    assert 1 not in s.eixos and 0 in s.eixos


def test_estado_tem_roi_e_eixo_espelhados_e_corte(sessao_pronta):
    t = comando(sessao_pronta, acao="estado")["estado"]["tempos"][0]
    assert t["roi_post"][0][0] == pytest.approx(64 - t["roi"][0][0])
    assert t["eixo_post"][0][0] == pytest.approx(64 - t["eixo"][0][0])
    seg = t["segmentacao"]
    assert t["eixo_origem"] == "automatico"
    assert seg["pixels_proximal"] == seg["pixels_distal"] == 144  # retângulo 12 × 24 cortado ao meio
    assert seg["ant"]["meio"] == pytest.approx([44, 28], abs=0.6) and seg["post"]["meio"] == pytest.approx([20, 28], abs=0.6)
    assert seg["ant"]["inicio"][1] < seg["ant"]["fim"][1]  # começa no topo


def test_eixo_automatico_e_correcao_manual(fantoma):
    s = Sessao()
    for n, d in fantoma.arquivos.items():
        s.adicionar_arquivo(n, d)
    assert "ROI" in comando(s, acao="estado")["estado"]["validacao"]["eixos"]
    s.desenhar_roi(0, fantoma.roi)
    e = comando(s, acao="estado")["estado"]
    assert e["validacao"]["eixos"] is None  # eixo automático em todos os tempos
    assert all(t["eixo_origem"] == "automatico" for t in e["tempos"])
    assert comando(s, acao="resultados")["ok"]
    assert not comando(s, acao="desenhar_eixo", tempo=0, pontos=[[40, 20], [40, 21]])["ok"]  # curto demais
    comando(s, acao="desenhar_eixo", tempo=3, pontos=[[44, 40], [44, 16]])  # traçado de baixo para cima
    assert s.eixos[3][0] == (44.0, 16.0)  # reorientado: começa no topo do fundo
    origens = [t["eixo_origem"] for t in comando(s, acao="estado")["estado"]["tempos"]]
    assert origens == ["automatico", "automatico", "automatico", "manual"]  # correção vale só no tempo
    comando(s, acao="copiar_eixo", tempo=3)
    assert all(t["eixo_origem"] == "manual" for t in comando(s, acao="estado")["estado"]["tempos"])
    comando(s, acao="eixo_automatico", tempo=1)
    assert comando(s, acao="estado")["estado"]["tempos"][1]["eixo_origem"] == "automatico"
    comando(s, acao="eixo_automatico")
    assert not s.eixos
    comando(s, acao="desenhar_eixo", tempo=2, pontos=[[44, 16], [44, 40]])
    comando(s, acao="desenhar_roi", tempo=2, pontos=fantoma.roi)  # ROI redesenhada descarta o eixo manual
    assert 2 not in s.eixos


def test_eixo_acompanha_o_arraste_da_roi(sessao_pronta, fantoma):
    def eixo(i):
        return comando(sessao_pronta, acao="estado")["estado"]["tempos"][i]["eixo"]

    antes = eixo(1)
    comando(sessao_pronta, acao="arrastar_roi", tempo=1, dx=2, dy=3, confirmar=True)
    assert np.allclose(np.array(eixo(1)), np.array(antes) + [2, 3])  # automático, recalculado
    sessao_pronta.desenhar_eixo(2, fantoma.eixo)
    comando(sessao_pronta, acao="arrastar_roi", tempo=2, dx=-1, dy=1, confirmar=True)
    assert sessao_pronta.eixos[2] == [(x - 1, y + 1) for x, y in fantoma.eixo]  # manual, transladado
    assert np.allclose(np.array(eixo(0)), np.array(antes))


def test_espacamento_de_pixel_da_comprimento_em_cm(fantoma):
    import pydicom

    s = Sessao()
    for n, d in fantoma.arquivos.items():
        ds = pydicom.dcmread(io.BytesIO(d))
        ds.PixelSpacing = [6.4, 6.4]
        b = io.BytesIO()
        ds.save_as(b)
        s.adicionar_arquivo(n, b.getvalue())
    s.desenhar_roi(0, fantoma.roi)
    s.desenhar_eixo(0, fantoma.eixo)  # eixo manual de 24 pixels
    seg = comando(s, acao="estado")["estado"]["tempos"][0]["segmentacao"]
    assert seg["comprimento_texto"] == "15,4 cm"  # 24 pixels × 6,4 mm
    assert seg["metade_texto"] == "7,7 cm"


def test_imagem_rgba(sessao_pronta):
    id_ = sessao_pronta.tempos[0].ant
    b = sessao_pronta.imagem({"arquivo": id_})
    assert len(b) == 64 * 64 * 4
    b2 = sessao_pronta.imagem(json.dumps({"arquivo": id_, "nivel": 10, "largura": 5, "invertido": True,
                                          "segmentacao": {"tempo": 0, "vista": "ant"}}))
    assert len(b2) == len(b) and b2 != b


def test_resultados_comando(sessao_pronta):
    r = comando(sessao_pronta, acao="resultados", decaimento=True, meia_vida="6,0067")
    assert r["ok"] and r["versao"] == __version__
    assert [l["textos"]["retencao"] for l in r["linhas"]] == ["100,0", "60,0", "30,0", "8,0"]
    assert [l["textos"]["razao_pd"] for l in r["linhas"]] == ["3,00", "1,22", "0,67", "0,43"]
    assert r["svg"].startswith("<svg") and "Proximal" in r["svg"]
    assert len(r["colunas"]) == 11
    assert r["referencia"] is None and all(l["referencia"] is None for l in r["linhas"])
    assert comando(sessao_pronta, acao="resultados", decaimento=True, meia_vida="abc")["ok"] is False


def test_resultados_com_referencia_adulto_e_pediatrica(sessao_pronta):
    r = comando(sessao_pronta, acao="resultados", faixa_etaria="adulto")
    assert r["referencia"]["grupo"] == "Adulto"
    assert [l["referencia"]["dentro"] if l["referencia"] else None for l in r["linhas"]] == [None, True, True, True]
    r = comando(sessao_pronta, acao="resultados", faixa_etaria="pediatrico", refeicao="ovo")
    assert "Cincinnati" in r["referencia"]["nota"]
    crit = [l["referencia"]["criterio"] if l["referencia"] else None for l in r["linhas"]]
    assert crit == [None, "esvaziamento > 10% e < 70%", "esvaziamento > 40%", "esvaziamento > 90%"]
    r = comando(sessao_pronta, acao="resultados", faixa_etaria="pediatrico", refeicao="aveia")
    assert [l["referencia"] is not None for l in r["linhas"]] == [False, True, False, False]
    assert not comando(sessao_pronta, acao="resultados", faixa_etaria="idoso")["ok"]


def test_resultados_bloqueados_sem_roi(fantoma):
    s = Sessao()
    for n, d in fantoma.arquivos.items():
        s.adicionar_arquivo(n, d)
    r = comando(s, acao="resultados")
    assert not r["ok"] and "ROI" in r["erro"]


def test_comando_desconhecido_e_invalido(sessao_pronta):
    assert not comando(sessao_pronta, acao="voar")["ok"]
    assert not comando(sessao_pronta, acao="desenhar_roi")["ok"]


def test_laudo_traz_versao_aviso_e_nao_traz_paciente(sessao_pronta):
    r = comando(sessao_pronta, acao="laudo", decaimento=True, meia_vida=6.0067, identificacao="Exame <teste>",
                versoes_extra={"pyodide": "0.29.5"}, faixa_etaria="pediatrico")
    html = r["html"]
    assert "Pediátrico" in html and "MacLean" in html and "Silver" in html
    assert f"mnp-nucleo {__version__}" in html and "Pyodide 0.29.5" in html
    assert AVISO in html
    assert "Exame &lt;teste&gt;" in html  # texto escapado
    assert "FANTOMA" not in html  # nome do paciente do DICOM não aparece


def test_exportar_csv(sessao_pronta):
    b = sessao_pronta.exportar({"formato": "csv", "decaimento": True, "meia_vida": 6.0067})
    texto = b.decode("utf-8-sig")
    linhas = list(csv.reader(io.StringIO(texto), delimiter=";"))
    assert linhas[0][0] == "Tempo" and len(linhas) == 5
    assert linhas[2][8] == "60,0" and linhas[2][6] == "1,22"
    b = sessao_pronta.exportar({"formato": "csv", "faixa_etaria": "adulto"})
    linhas = list(csv.reader(io.StringIO(b.decode("utf-8-sig")), delimiter=";"))
    assert linhas[0][-2:] == ["Referência (adulto)", "Situação"] and linhas[2][-1] == "dentro da referência"


def test_exportar_xlsx(sessao_pronta):
    b = sessao_pronta.exportar({"formato": "xlsx", "decaimento": True, "meia_vida": 6.0067})
    wb = openpyxl.load_workbook(io.BytesIO(b))
    ws = wb.active
    assert ws["A1"].value.startswith("Esvaziamento")
    assert "não informada" in ws["A3"].value
    assert __version__ in ws["A4"].value
    assert ws["A6"].value == "Tempo"
    assert ws["I8"].value == pytest.approx(60.0)
    assert ws["G7"].value == pytest.approx(3.0)
    assert ws["C7"].value == round(sessao_pronta.resultados()[0].mg_total)


def test_protocolo_e_lote_reproduzem_o_navegador(sessao_pronta, fantoma, tmp_path):
    comando(sessao_pronta, acao="arrastar_roi", tempo=3, dx=1, dy=1, confirmar=True)
    esperado = [l.retencao for l in sessao_pronta.resultados(True, 6.0067)]
    protocolo = comando(sessao_pronta, acao="protocolo", decaimento=True, meia_vida=6.0067)["protocolo"]
    assert protocolo["versao_nucleo"] == __version__
    pasta = tmp_path / "dicoms"
    pasta.mkdir()
    for nome, dados in fantoma.arquivos.items():
        (pasta / ("renomeado_" + nome)).write_bytes(dados)  # acha pelo SHA-256 mesmo com outro nome
    _, linhas, _ = lote.processar_protocolo(protocolo, pasta)
    assert [l.retencao for l in linhas] == esperado

    arq = tmp_path / "p.json"
    arq.write_text(json.dumps(protocolo), encoding="utf-8")
    assert lote.main([str(arq), "--pasta", str(pasta), "--csv", str(tmp_path / "r.csv"),
                      "--laudo", str(tmp_path / "l.html"), "--faixa-etaria", "adulto"]) == 0
    assert (tmp_path / "r.csv").read_text().count("\n") == 5
    laudo = (tmp_path / "l.html").read_text(encoding="utf-8")
    assert __version__ in laudo and "Adulto" in laudo and "Tougas" in laudo


def test_protocolo_antigo_recusado(sessao_pronta, tmp_path):
    from mnp_nucleo.gastrico.sessao import ErroSessao

    protocolo = sessao_pronta.protocolo()
    protocolo["formato"] = 1
    with pytest.raises(ErroSessao, match="0.1"):
        lote.processar_protocolo(protocolo, tmp_path)


def test_lote_detecta_arquivo_diferente(sessao_pronta, fantoma, tmp_path):
    protocolo = sessao_pronta.protocolo()
    for nome, dados in fantoma.arquivos.items():
        if nome == "T1_ant.dcm":
            dados = dados[:-1] + bytes([(dados[-1] + 1) % 256])  # um pixel alterado
        (tmp_path / nome).write_bytes(dados)
    from mnp_nucleo.gastrico.sessao import ErroSessao

    with pytest.raises(ErroSessao, match="não encontrado"):
        lote.processar_protocolo(protocolo, tmp_path)


def test_carregar_fantoma_pela_interface():
    s = Sessao()
    r = comando(s, acao="carregar_fantoma")
    assert r["ok"] and len(r["estado"]["tempos"]) == 4
    comando(s, acao="desenhar_roi", tempo=0, pontos=r["fantoma"]["roi"])
    comando(s, acao="desenhar_eixo", tempo=0, pontos=r["fantoma"]["eixo"])
    res = comando(s, acao="resultados", decaimento=True, meia_vida=6.0067)
    assert [l["textos"]["retencao"] for l in res["linhas"]] == ["100,0", "60,0", "30,0", "8,0"]
    r = comando(s, acao="carregar_fantoma", multiframe=True)
    assert len(r["estado"]["sugestoes"]) == 4


def test_traco_fora_da_imagem_e_trazido_para_a_borda(fantoma):
    s = Sessao()
    for n, d in fantoma.arquivos.items():
        s.adicionar_arquivo(n, d)
    comando(s, acao="desenhar_roi", tempo=0, pontos=[[-10, -10], [80, -5], [80, 90], [-3, 90]])
    xs = [p[0] for p in s.rois[0]]
    ys = [p[1] for p in s.rois[0]]
    assert min(xs) == 0 and max(xs) == 64 and min(ys) == 0 and max(ys) == 64
