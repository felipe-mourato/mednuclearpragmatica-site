"""Testes da cintilografia renal com DMSA (núcleo mnp_nucleo.dmsa).

Valores esperados calculados à mão a partir do desenho dos fantomas.
"""

import io
import json
import math
import struct
import zlib

import numpy as np
import pytest

from mnp_nucleo.comum import morfologia
from mnp_nucleo.comum.png import codificar_png
from mnp_nucleo.comum.roi import mascara_roi
from mnp_nucleo.dmsa import calculo, figuras, incidencias, referencias
from mnp_nucleo.dmsa.fundo import anel_perirrenal
from mnp_nucleo.dmsa.lote import main as lote_main
from mnp_nucleo.dmsa.sessao import Sessao
from mnp_nucleo.fantomas import fantoma_dmsa, gerar_dicom


def cmd(s, **kw):
    return json.loads(s.executar(kw))


@pytest.fixture
def fd():
    return fantoma_dmsa()


@pytest.fixture
def sessao(fd):
    s = Sessao()
    for nome, dados in fd.arquivos.items():
        s.adicionar_arquivo(nome, dados)
    for rim in "DE":
        assert cmd(s, acao="desenhar_roi", rim=rim, pontos=fd.rois[rim])["ok"]
    return s


def resultado(s, chave):
    r = cmd(s, acao="resultados")
    assert r["ok"], r
    return next(l for l in r["funcao"] if l["chave"] == chave)


# ------------------------------------------------------------- incidências
@pytest.mark.parametrize(
    "descricao, esperado",
    [
        ("DMSA POST", "POST"),
        ("POSTERIOR", "POST"),
        ("Renal PÓS", "POST"),
        ("DMSA ANT", "ANT"),
        ("ANTERIOR", "ANT"),
        ("OPD", "OPD"),
        ("RPO 3h", "OPD"),
        ("LPO", "OPE"),
        ("OAD", "OAD"),
        ("LAO", "OAE"),
        ("ANT/POST", "POST"),
        ("", "POST"),
        (None, "POST"),
        ("QUANTITATIVO", "POST"),
    ],
)
def test_adivinhar_incidencia(descricao, esperado):
    assert incidencias.adivinhar_incidencia(descricao)[0] == esperado


def test_obliqua_sem_lado_vem_com_aviso():
    inc, aviso = incidencias.adivinhar_incidencia("OBL")
    assert inc == "OPD" and "confira" in aviso


# ---------------------------------------------------------------- morfologia
def test_disco_raio_4_tem_49_deslocamentos():
    # contagem à mão de dy² + dx² ≤ 16: 1 + 4×4 (eixos) + 4×8 (quadrantes) = 49
    assert len(morfologia.deslocamentos_disco(4)) == 49


def test_anel_de_retangulo_tem_area_calculada_a_mao():
    # Retângulo 10 × 22: faixas laterais 4 × (2·10 + 2·22) = 256 e 8 pixels por canto
    # (dx=1: 3, dx=2: 3, dx=3: 2) → 288.
    m = np.zeros((64, 64), dtype=bool)
    m[18:40, 38:48] = True
    anel = anel_perirrenal(m, np.zeros_like(m), 4)
    assert anel.sum() == 288
    assert not (anel & m).any()


def test_anel_nao_invade_o_outro_rim_e_para_na_borda():
    m = np.zeros((40, 40), dtype=bool)
    m[10:20, 0:5] = True  # encostado na borda esquerda
    outro = np.zeros_like(m)
    outro[10:20, 7:12] = True  # a 2 pixels de distância
    anel = anel_perirrenal(m, outro, 4)
    assert not (anel & outro).any()
    assert not (anel & m).any()
    # nenhum pixel do anel além de 4 pixels do rim
    ys, xs = np.nonzero(anel)
    assert xs.max() <= 4 + 4


def test_rle_ida_e_volta():
    rng = np.random.default_rng(1)
    m = rng.random((17, 23)) > 0.6
    assert np.array_equal(morfologia.rle_decodificar(morfologia.rle_codificar(m), 17, 23), m)
    assert morfologia.rle_codificar(np.zeros((3, 3), dtype=bool)) == []


def test_rle_fora_da_imagem_falha():
    with pytest.raises(ValueError):
        morfologia.rle_decodificar([[5, 10]], 3, 3)


# ----------------------------------------------------------------------- PNG
def test_png_decodifica_para_os_mesmos_pixels():
    rgba = np.zeros((3, 2, 4), dtype=np.uint8)
    rgba[..., 0] = [[1, 2], [3, 4], [5, 6]]
    rgba[..., 3] = 255
    png = codificar_png(rgba)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    largura, altura = struct.unpack(">II", png[16:24])
    assert (largura, altura) == (2, 3)
    i = png.index(b"IDAT")
    n = struct.unpack(">I", png[i - 4 : i])[0]
    cru = zlib.decompress(png[i + 4 : i + 4 + n])
    linhas = np.frombuffer(cru, dtype=np.uint8).reshape(3, 1 + 2 * 4)
    assert (linhas[:, 0] == 0).all()
    assert np.array_equal(linhas[:, 1:].reshape(3, 2, 4), rgba)


# ------------------------------------------------------------------- cálculo
def test_contagem_corrigida_a_mao():
    # 24200 − 10 × 220 = 22000
    assert calculo.contagem_corrigida(24200, 220, 10.0) == 22000
    # fundo maior que o rim: zero, nunca negativo
    assert calculo.contagem_corrigida(100, 50, 5.0) == 0.0


def test_funcao_relativa_e_par_somam_100():
    d = calculo.funcao_relativa(22000, 28000)
    assert d == pytest.approx(44.0)
    pd, pe = calculo.arredondar_par(44.96)
    assert (pd, pe) == (45.0, 55.0)
    assert pd + pe == 100.0
    pd, pe = calculo.arredondar_par(100 / 3)
    assert (pd, pe) == (33.3, 66.7)


def test_correcao_pelo_tamanho_exemplo_dos_rins_limitrofes():
    # Rins de 44% e 56% do tamanho com a mesma função por pixel: FRR 44%, corrigida 50%.
    assert calculo.funcao_relativa(44 * 7, 56 * 7) == pytest.approx(44.0)
    assert calculo.funcao_corrigida_tamanho(44 * 7, 44, 56 * 7, 56) == pytest.approx(50.0)


def test_media_geometrica_a_mao():
    # √(22000 × 11000) = 15556,349…; √(28000 × 11200) = 17708,754…
    d = {
        "D": calculo.ContagensRim("D", "post", 0, 220, 0, 0, 0, 22000.0),
        "E": calculo.ContagensRim("E", "post", 0, 280, 0, 0, 0, 28000.0),
    }
    a = {
        "D": calculo.ContagensRim("D", "ant", 0, 220, 0, 0, 0, 11000.0),
        "E": calculo.ContagensRim("E", "ant", 0, 280, 0, 0, 0, 11200.0),
    }
    r = calculo.calcular(d, a)
    assert r.mg_contagens["D"] == pytest.approx(15556.349, abs=1e-3)
    assert r.mg_contagens["E"] == pytest.approx(17708.754, abs=1e-3)
    assert r.parametro("frr_mg").d == pytest.approx(46.7648, abs=1e-4)
    # contagem por pixel: √(50 × 100) = 70,711 e √(40 × 100) = 63,246 → 52,786%
    assert r.parametro("tam_mg").d == pytest.approx(52.7864, abs=1e-4)


def test_sem_anterior_sem_media_geometrica():
    d = {
        "D": calculo.ContagensRim("D", "post", 0, 10, 0, 0, 0, 30.0),
        "E": calculo.ContagensRim("E", "post", 0, 10, 0, 0, 0, 70.0),
    }
    r = calculo.calcular(d)
    assert r.parametro("frr_post").d_arred == 30.0
    assert not r.parametro("frr_mg").disponivel


def test_soma_zero_falha():
    with pytest.raises(ValueError):
        calculo.funcao_relativa(0, 0)


@pytest.mark.parametrize("valor, dentro", [(45.0, True), (55.0, True), (50.0, True), (44.9, False), (55.1, False)])
def test_faixa_normal_limites_incluidos(valor, dentro):
    assert referencias.comparar_funcao(valor)["dentro"] is dentro


# ------------------------------------------------------- sessão (fantoma)
CHAVES = ("frr_post", "frr_mg", "area_post", "area_mg", "tam_post", "tam_mg")


def test_fantoma_de_ponta_a_ponta(sessao, fd):
    for chave in CHAVES:
        l = resultado(sessao, chave)
        assert l["d"] == pytest.approx(fd.esperado[chave], abs=1e-9), chave
        assert l["d_arred"] + l["e_arred"] == pytest.approx(100.0)
    assert resultado(sessao, "frr_post")["texto_d"] == "44,0"
    assert resultado(sessao, "frr_post")["referencia"]["dentro"] is False
    assert resultado(sessao, "tam_post")["texto_d"] == "50,0"
    assert resultado(sessao, "tam_post")["referencia"] is None
    assert resultado(sessao, "frr_mg")["texto_d"] == "46,8"


def test_fantoma_contagens(sessao, fd):
    r = cmd(sessao, acao="resultados")
    linhas = {(l["rim"], l["vista"]): l["valores"] for l in r["contagens"]["linhas"]}
    assert linhas[("D", "post")]["bruta"] == 220 * 110
    assert linhas[("D", "post")]["fundo_por_pixel"] == 10.0
    assert linhas[("D", "post")]["fundo_pixels"] == 288
    assert linhas[("E", "post")]["fundo_pixels"] == 336
    assert linhas[("D", "post")]["corrigida"] == fd.contagens["D_post"]
    assert linhas[("E", "ant")]["corrigida"] == fd.contagens["E_ant"]


def test_multiframe_quadro_1_posterior_quadro_2_anterior(fd):
    s = Sessao()
    s.adicionar_arquivo("mf.dcm", fd.multiframe["dmsa_post_ant.dcm"])
    e = cmd(s, acao="estado")["estado"]
    assert [x["incidencia"] for x in e["slots"]] == ["POST", "ANT"]
    assert [x["rotulo"] for x in e["slots"]] == ["mf.dcm · quadro 1", "mf.dcm · quadro 2"]
    for rim in "DE":
        cmd(s, acao="desenhar_roi", rim=rim, pontos=fd.rois[rim])
    for chave in CHAVES:
        assert resultado(s, chave)["d"] == pytest.approx(fd.esperado[chave])


def test_rescale_aplicado(fd):
    # Mesma imagem com pixels dobrados e RescaleSlope 0,5: mesmo resultado.
    s = Sessao()
    base = fantoma_dmsa()
    s.adicionar_arquivo("post.dcm", gerar_dicom(_pixels(base, "dmsa_post.dcm") * 2, None, "POST", inclinacao=0.5))
    s.adicionar_arquivo("ant.dcm", gerar_dicom(_pixels(base, "dmsa_ant.dcm") * 2, None, "ANT", inclinacao=0.5))
    for rim in "DE":
        cmd(s, acao="desenhar_roi", rim=rim, pontos=fd.rois[rim])
    for chave in CHAVES:
        assert resultado(s, chave)["d"] == pytest.approx(fd.esperado[chave])


def _pixels(f, nome):
    from mnp_nucleo.comum.dicom import ler_dicom

    return ler_dicom(f.arquivos[nome]).quadro(0).astype(np.int64)


def test_espelhamento_preserva_area(sessao):
    for rim in "DE":
        assert sessao._mascara_rim("post", rim).sum() == sessao._mascara_rim("ant", rim).sum()
    # D em x 38..48 na posterior → 16..26 na anterior
    xs = [p[0] for p in sessao.roi_vista("ant", "D")]
    assert (min(xs), max(xs)) == (16.0, 26.0)


def test_exibicao_nao_muda_contagens(sessao):
    antes = cmd(sessao, acao="resultados")["funcao"]
    e = cmd(sessao, acao="estado")["estado"]
    sessao.imagem({"slot": e["post"], "nivel": 3, "largura": 2, "invertido": True, "fundo": True})
    sessao.exportar({"formato": "png", "slot": e["post"], "invertido": True, "janela": {"nivel": 1, "largura": 1}})
    assert cmd(sessao, acao="resultados")["funcao"] == antes


def test_sem_fundo(sessao):
    assert cmd(sessao, acao="definir_fundo", metodo="nenhum")["ok"]
    # sem subtração: 220 × 110 ÷ (220 × 110 + 280 × 110) = 44,0%; por pixel 110 ÷ 220 = 50%
    assert resultado(sessao, "frr_post")["d"] == pytest.approx(44.0)
    assert resultado(sessao, "tam_post")["d"] == pytest.approx(50.0)
    # MG sem fundo: √(24200 × 12760) ÷ (… + √(30800 × 13440))
    mg_d = math.sqrt(24200 * 12760)
    mg_e = math.sqrt(30800 * 13440)
    assert resultado(sessao, "frr_mg")["d"] == pytest.approx(100 * mg_d / (mg_d + mg_e))


def test_sem_anterior_media_geometrica_aparece_como_traco(fd):
    s = Sessao()
    s.adicionar_arquivo("post.dcm", fd.arquivos["dmsa_post.dcm"])
    for rim in "DE":
        cmd(s, acao="desenhar_roi", rim=rim, pontos=fd.rois[rim])
    l = resultado(s, "frr_mg")
    assert l["texto_d"] == "—" and l["referencia"] is None
    assert resultado(s, "frr_post")["d"] == pytest.approx(44.0)


# ------------------------------------------------------------- validações
def test_sem_posterior_nao_avanca(fd):
    s = Sessao()
    s.adicionar_arquivo("ant.dcm", fd.arquivos["dmsa_ant.dcm"])
    e = cmd(s, acao="estado")["estado"]
    assert "POSTERIOR" in e["validacao"]["arquivos"]


def test_posterior_duplicada_bloqueia(fd):
    s = Sessao()
    s.adicionar_arquivo("a.dcm", fd.arquivos["dmsa_post.dcm"])
    s.adicionar_arquivo("b.dcm", fd.arquivos["dmsa_post.dcm"])
    e = cmd(s, acao="estado")["estado"]
    assert "2 imagens" in e["validacao"]["arquivos"]
    assert e["duplicadas"] == ["POST"]
    cmd(s, acao="definir_incidencia", slot=e["slots"][1]["id"], incidencia="IGNORAR")
    assert cmd(s, acao="estado")["estado"]["validacao"]["arquivos"] is None


def test_obliqua_duplicada_so_avisa(fd, sessao):
    e = cmd(sessao, acao="estado")["estado"]
    ope = next(x for x in e["slots"] if x["incidencia"] == "OPE")
    e = cmd(sessao, acao="definir_incidencia", slot=ope["id"], incidencia="OPD")["estado"]
    assert e["validacao"]["fundo"] is None
    assert any("OPD" in a for a in e["avisos"])


def test_sem_as_duas_rois_nao_avanca(fd):
    s = Sessao()
    s.adicionar_arquivo("post.dcm", fd.arquivos["dmsa_post.dcm"])
    cmd(s, acao="desenhar_roi", rim="D", pontos=fd.rois["D"])
    e = cmd(s, acao="estado")["estado"]
    assert "rim esquerdo" in e["validacao"]["rois"]
    assert not cmd(s, acao="resultados")["ok"]


def test_rois_sobrepostas_bloqueiam(fd):
    s = Sessao()
    s.adicionar_arquivo("post.dcm", fd.arquivos["dmsa_post.dcm"])
    cmd(s, acao="desenhar_roi", rim="D", pontos=fd.rois["D"])
    cmd(s, acao="desenhar_roi", rim="E", pontos=fd.rois["D"])
    assert "sobrepõem" in cmd(s, acao="estado")["estado"]["validacao"]["rois"]


def test_aviso_rins_trocados(fd):
    s = Sessao()
    s.adicionar_arquivo("post.dcm", fd.arquivos["dmsa_post.dcm"])
    cmd(s, acao="desenhar_roi", rim="D", pontos=fd.rois["E"])
    cmd(s, acao="desenhar_roi", rim="E", pontos=fd.rois["D"])
    assert any("trocados" in a for a in cmd(s, acao="estado")["estado"]["avisos"])


def test_dimensoes_diferentes_avisa_e_escala(fd):
    s = Sessao()
    s.adicionar_arquivo("post.dcm", fd.arquivos["dmsa_post.dcm"])
    ant = np.kron(_pixels(fd, "dmsa_ant.dcm"), np.ones((2, 2), dtype=np.int64))  # 128 × 128
    s.adicionar_arquivo("ant.dcm", gerar_dicom(ant, None, "ANT"))
    for rim in "DE":
        cmd(s, acao="desenhar_roi", rim=rim, pontos=fd.rois[rim])
    e = cmd(s, acao="estado")["estado"]
    assert any("dimensões" in a for a in e["avisos"])
    xs = [p[0] for p in s.roi_vista("ant", "D")]
    assert (min(xs), max(xs)) == (32.0, 52.0)
    assert s._mascara_rim("ant", "D").sum() == 4 * 220


def test_arquivo_invalido_vira_mensagem():
    s = Sessao()
    e = json.loads(s.adicionar_arquivo("x.txt", b"nao e dicom" * 30))["estado"]
    assert e["erros"] and "x.txt" in e["erros"][0]


# ------------------------------------------------------------------- arraste
def test_arrastar_posterior_move_anterior_espelhada(sessao):
    antes = sessao.roi_vista("ant", "D")
    r = cmd(sessao, acao="arrastar_roi", vista="post", rim="D", dx=2, dy=1, confirmar=True)
    assert r["ok"]
    depois = sessao.roi_vista("ant", "D")
    assert depois[0][0] == antes[0][0] - 2 and depois[0][1] == antes[0][1] + 1


def test_arrastar_anterior_afeta_so_a_anterior(sessao, fd):
    post = sessao.roi_vista("post", "E")
    cmd(sessao, acao="arrastar_roi", vista="ant", rim="E", dx=3, dy=0, confirmar=True)
    assert sessao.roi_vista("post", "E") == post
    e = cmd(sessao, acao="estado")["estado"]
    assert e["deslocada"]["E"] is True
    cmd(sessao, acao="restaurar_espelho", rim="E")
    assert cmd(sessao, acao="estado")["estado"]["deslocada"]["E"] is False


def test_arraste_limitado_a_imagem(sessao):
    cmd(sessao, acao="arrastar_roi", vista="post", rim="D", dx=500, dy=-500, confirmar=True)
    xs = [p[0] for p in sessao.roi_vista("post", "D")]
    ys = [p[1] for p in sessao.roi_vista("post", "D")]
    assert max(xs) == 64 and min(ys) == 0


def test_arraste_sem_confirmar_nao_grava(sessao):
    antes = sessao.roi_vista("post", "D")
    r = cmd(sessao, acao="arrastar_roi", vista="post", rim="D", dx=2, dy=0, confirmar=False)
    assert "roi_ant" in r
    assert sessao.roi_vista("post", "D") == antes


def test_toca_roi(sessao):
    assert cmd(sessao, acao="toca_roi", vista="post", x=40, y=25)["rim"] == "D"
    assert cmd(sessao, acao="toca_roi", vista="ant", x=42, y=25)["rim"] == "E"
    assert cmd(sessao, acao="toca_roi", vista="post", x=2, y=2)["rim"] is None


# -------------------------------------------------------------- fundo à mão
def test_fundo_desenhado_a_mao_e_volta_ao_automatico(sessao):
    # Quadrado 6 × 6 fora dos rins, na posterior, para o rim direito: fundo 10 por pixel.
    quad = [(52, 50), (58, 50), (58, 56), (52, 56)]
    e = cmd(sessao, acao="desenhar_fundo", vista="post", rim="D", pontos=quad)["estado"]
    reg = e["fundo"]["regioes"]["post"]["D"]
    assert reg["origem"] == "manual" and reg["pixels"] == 36
    assert resultado(sessao, "frr_post")["d"] == pytest.approx(44.0)
    e = cmd(sessao, acao="fundo_automatico", vista="post", rim="D")["estado"]
    assert e["fundo"]["regioes"]["post"]["D"] == {"pixels": 288, "origem": "automatico", "poligono": None}


def test_fundo_manual_dentro_do_rim_e_recusado(sessao, fd):
    r = cmd(sessao, acao="desenhar_fundo", vista="post", rim="D", pontos=fd.rois["D"])
    assert not r["ok"] and "fora dos rins" in r["erro"]


def test_redesenhar_roi_descarta_fundo_manual(sessao, fd):
    cmd(sessao, acao="desenhar_fundo", vista="post", rim="D", pontos=[(52, 50), (58, 50), (58, 56), (52, 56)])
    cmd(sessao, acao="desenhar_roi", rim="D", pontos=fd.rois["D"])
    assert cmd(sessao, acao="estado")["estado"]["fundo"]["regioes"]["post"]["D"]["origem"] == "automatico"


# ------------------------------------------------------- saídas e protocolo
def test_laudo_sem_dados_do_dicom_com_identificacao_digitada(sessao):
    r = cmd(sessao, acao="laudo", paciente="Maria Teste", data_exame="13/10/2026", versoes_extra={"pyodide": "0.29.5"})
    html = r["html"]
    assert "FANTOMA" not in html and "SINTETICO" not in html  # PatientName do DICOM
    assert "Maria Teste" in html and "13/10/2026" in html
    assert "ainda não validada para uso clínico" in html
    assert "10.1016/j.ejmp.2022.01.010" in html and "10.1007/s40336-022-00484-x" in html
    assert "10.1097/00006231-200402000-00013" in html
    assert "Pyodide 0.29.5" in html
    assert html.count("data:image/png;base64,") == 4  # POST, ANT, OPD, OPE
    assert html.index("Posterior") < html.index("Anterior")


def test_laudo_sem_identificacao_nao_tem_campos(sessao):
    html = cmd(sessao, acao="laudo")["html"]
    assert "Data do exame" not in html


def test_figura_png_tem_tamanho_esperado(sessao):
    e = cmd(sessao, acao="estado")["estado"]
    png = sessao.exportar({"formato": "png", "slot": e["post"]})
    assert struct.unpack(">II", png[16:24]) == (384, 384)  # 64 × 6
    rgba = figuras.quadro_rgba(np.zeros((40, 20)), 0.5, 1.0)
    assert rgba.shape == (400, 400, 4)  # lado = 40 × 10, imagem centralizada


def test_csv_e_xlsx(sessao):
    from openpyxl import load_workbook

    texto = sessao.exportar({"formato": "csv"}).decode("utf-8-sig")
    assert "Função relativa, posterior;44,0;56,0" in texto
    wb = load_workbook(io.BytesIO(sessao.exportar({"formato": "xlsx"})))
    valores = [c for linha in wb.active.iter_rows(values_only=True) for c in linha if c is not None]
    assert 44.0 in valores and 50.0 in valores
    assert any("Civan" in str(v) for v in valores)


def test_tabelas_para_copiar(sessao):
    r = cmd(sessao, acao="resultados")
    assert "<table" in r["tabelas_html"]
    assert "Função relativa, posterior\t44,0\t56,0\tfora" in r["tabelas_texto"]


def test_protocolo_reproduz_resultado(sessao, fd, tmp_path):
    cmd(sessao, acao="arrastar_roi", vista="ant", rim="D", dx=1, dy=0, confirmar=True)
    cmd(sessao, acao="desenhar_fundo", vista="ant", rim="E", pontos=[(2, 2), (10, 2), (10, 10), (2, 10)])
    esperado = cmd(sessao, acao="resultados")["funcao"]
    protocolo = json.loads(sessao.exportar({"formato": "protocolo"}))
    assert protocolo["fundo"]["regioes"]["post"]["D"]["origem"] == "automatico"
    assert protocolo["fundo"]["regioes"]["ant"]["E"]["origem"] == "manual"
    for nome, dados in fd.arquivos.items():
        (tmp_path / nome).write_bytes(dados)
    arquivos = {nome: dados for nome, dados in fd.arquivos.items()}

    def ler(nome, sha):
        return arquivos[nome]

    s2 = Sessao.de_protocolo(protocolo, ler)
    assert cmd(s2, acao="resultados")["funcao"] == esperado

    caminho = tmp_path / "exame.json"
    caminho.write_text(json.dumps(protocolo), encoding="utf-8")
    assert lote_main([str(caminho), "--pasta", str(tmp_path), "--csv", str(tmp_path / "r.csv"), "--laudo", str(tmp_path / "l.html")]) == 0
    csv_texto = (tmp_path / "r.csv").read_text(encoding="utf-8")
    d = next(l for l in esperado if l["chave"] == "frr_post")["d_arred"]
    assert f"{d}" in csv_texto
    assert "Laudo de processamento" in (tmp_path / "l.html").read_text(encoding="utf-8")


def test_protocolo_usa_mascara_gravada(sessao, fd):
    protocolo = json.loads(sessao.exportar({"formato": "protocolo"}))
    # Troca o anel do rim direito na posterior por uma máscara de 1 pixel com fundo 10:
    protocolo["fundo"]["regioes"]["post"]["D"]["mascara_rle"] = [[0, 1]]
    s2 = Sessao.de_protocolo(protocolo, lambda nome, sha: fd.arquivos[nome])
    r = cmd(s2, acao="resultados")
    linha = next(l for l in r["contagens"]["linhas"] if l["rim"] == "D" and l["vista"] == "post")
    assert linha["valores"]["fundo_pixels"] == 1


def test_protocolo_com_arquivo_trocado_falha(sessao, fd):
    protocolo = json.loads(sessao.exportar({"formato": "protocolo"}))
    with pytest.raises(Exception, match="SHA-256"):
        Sessao.de_protocolo(protocolo, lambda nome, sha: fd.arquivos["dmsa_ant.dcm"])


def test_protocolo_de_outra_ferramenta():
    with pytest.raises(Exception, match="outra ferramenta"):
        Sessao.de_protocolo({"ferramenta": "esvaziamento-gastrico", "imagens": []}, lambda n, s: b"")


def test_carregar_fantoma_pela_sessao():
    s = Sessao()
    r = cmd(s, acao="carregar_fantoma")
    assert r["ok"] and len(r["estado"]["slots"]) == 4
    assert r["fantoma"]["esperado"]["frr_post"] == pytest.approx(44.0)
    r = cmd(s, acao="carregar_fantoma", multiframe=True)
    assert [x["incidencia"] for x in r["estado"]["slots"]] == ["POST", "ANT"]


def test_comando_desconhecido_e_invalido(sessao):
    assert not cmd(sessao, acao="xyz")["ok"]
    assert not cmd(sessao, acao="desenhar_roi", rim="X", pontos=[])["ok"]
    assert not cmd(sessao, acao="desenhar_roi", rim="D", pontos=[[0, 0], [1, 0], [1, 1]])["ok"]


def test_mascara_espelhada_igual_a_mascara_do_fantoma(sessao, fd):
    m = sessao._mascara_rim("ant", "E")
    esperado = mascara_roi([(64 - x, y) for x, y in fd.rois["E"]], 64, 64)
    assert np.array_equal(m, esperado)
